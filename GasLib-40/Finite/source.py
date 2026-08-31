#!/usr/bin/env python3
"""Update source-pressure rows in this folder's input workbooks."""

from __future__ import annotations

import argparse
import shutil
from pathlib import Path

from openpyxl import load_workbook


BASE_DIR = Path(__file__).resolve().parent
INPUT_DIR = BASE_DIR / "Input_data"
DEFAULT_FILES = (
    INPUT_DIR / "inputData.xlsx",
    INPUT_DIR / "inputData_longer_horizon.xlsx",
)


def contained_in_base(path: Path) -> Path:
    resolved = path.resolve()
    try:
        resolved.relative_to(BASE_DIR)
    except ValueError as exc:
        raise ValueError(
            f"Refusing to access a file outside {BASE_DIR}: {resolved}"
        ) from exc
    return resolved


def update_workbook(source: Path, output: Path, pressure_pa: float, backup: bool) -> None:
    source = contained_in_base(source)
    output = contained_in_base(output)
    if not source.is_file():
        raise FileNotFoundError(source)

    if backup and source == output:
        backup_path = source.with_suffix(source.suffix + ".bak")
        shutil.copy2(source, backup_path)
        print(f"Backup: {backup_path}")

    workbook = load_workbook(source)
    if "SourcesSP" not in workbook.sheetnames:
        raise ValueError(f"SourcesSP sheet missing in {source}")

    sheet = workbook["SourcesSP"]
    sources = []
    changed = 0
    for row in range(2, sheet.max_row + 1):
        source_name = sheet.cell(row, 1).value
        variable = sheet.cell(row, 2).value
        if not str(source_name).startswith("source_") or str(variable).strip().lower() != "p":
            continue
        sources.append(str(source_name))
        for column in range(3, sheet.max_column + 1):
            sheet.cell(row, column).value = pressure_pa
            changed += 1

    if not sources:
        raise ValueError(f"No source pressure rows found in {source}")

    output.parent.mkdir(parents=True, exist_ok=True)
    workbook.save(output)
    print(f"Output: {output}")
    print(f"Sources: {', '.join(sources)}")
    print(f"Changed cells: {changed}")


def parse_args() -> argparse.Namespace:
    parser = argparse.ArgumentParser(
        description=(
            "Set source pressure in both Input_data workbooks by default. "
            "The pressure unit is Pa."
        )
    )
    parser.add_argument("--pressure-pa", required=True, type=float)
    parser.add_argument(
        "--in-place",
        action="store_true",
        help="overwrite both input workbooks and create .bak backups",
    )
    return parser.parse_args()


def main() -> None:
    args = parse_args()
    if not (1.0e5 <= args.pressure_pa <= 1.0e8):
        raise ValueError(
            "--pressure-pa must be between 100,000 and 100,000,000 Pa; "
            "check that the supplied unit is Pa"
        )
    if not args.in_place:
        raise ValueError(
            "Use --in-place to confirm modifying both workbooks in Input_data"
        )

    for path in DEFAULT_FILES:
        update_workbook(path, path, args.pressure_pa, backup=True)

    print(
        f"Pressure: {args.pressure_pa:,.3f} Pa "
        f"({args.pressure_pa / 1.0e5:.5f} bar)"
    )


if __name__ == "__main__":
    main()

# cd /home/yumings/Gas_pipeline_inf_enmpc/GasLib-40_finite_backup_1
# python source.py --pressure-pa 3501325 --in-place